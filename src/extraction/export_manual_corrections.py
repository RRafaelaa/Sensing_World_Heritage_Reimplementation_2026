#!/usr/bin/env python3
"""Export cell-level corrections between two review workbooks.

The output is deliberately flat and path-free so that manual corrections can
be replayed without publishing the reviewers' original local workbooks.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("before", type=Path)
    parser.add_argument("after", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--sheet", default="Gemini_2_5_pro")
    parser.add_argument("--key-column", default="filename")
    args = parser.parse_args()

    before = load_workbook(args.before, read_only=False, data_only=False)
    after = load_workbook(args.after, read_only=False, data_only=False)
    old_sheet = before[args.sheet]
    new_sheet = after[args.sheet]

    headers = [cell.value for cell in old_sheet[1]]
    if args.key_column not in headers:
        raise ValueError(f"Missing key column: {args.key_column}")
    key_index = headers.index(args.key_column) + 1

    corrections: list[dict[str, object]] = []
    for row_index in range(2, max(old_sheet.max_row, new_sheet.max_row) + 1):
        record_id = new_sheet.cell(row_index, key_index).value
        for column_index in range(1, max(old_sheet.max_column, new_sheet.max_column) + 1):
            old_value = old_sheet.cell(row_index, column_index).value
            new_value = new_sheet.cell(row_index, column_index).value
            if old_value == new_value:
                continue
            field = headers[column_index - 1] if column_index <= len(headers) else None
            corrections.append(
                {
                    "sheet": args.sheet,
                    "row_number": row_index,
                    "record_id": record_id,
                    "field": field,
                    "old_value": old_value,
                    "new_value": new_value,
                }
            )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "sheet",
                "row_number",
                "record_id",
                "field",
                "old_value",
                "new_value",
            ],
        )
        writer.writeheader()
        writer.writerows(corrections)

    print(f"Exported {len(corrections)} corrections to {args.output}")


if __name__ == "__main__":
    main()
